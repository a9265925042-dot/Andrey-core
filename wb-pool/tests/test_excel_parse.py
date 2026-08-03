"""Excel sheet parsers против синтетического workbook (offline).

Заголовки соответствуют документированному формату WB экспорта
(docs/skill/docs/06-excel-ingest.md). При изменении реального формата WB
эти тесты фиксируют контракт парсеров, не сам WB.
"""

import io

from openpyxl import Workbook

from wb_pool.etl.excel import (
    parse_search_queries_sheet,
    parse_size_stocks_sheet,
    parse_warehouse_sheet,
    read_workbook,
)


def _make_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Поисковые запросы"
    ws.append(
        [
            "Поисковый запрос",
            "Артикул WB",
            "Частота",
            "Динамика частоты",
            "Переходы в корзину",
            "Заказы",
            "Конверсия в корзину, %",
            "Конверсия в заказ, %",
        ]
    )
    ws.append(["крем для ног", 111, 5000, 250, 40, 7, 10.0, 2.0])
    ws.append(["крем для ног", 222, 5000, 250, 15, 3, 5.0, 1.0])
    ws.append(["воск от трещин", 111, 900, -30, 12, 2, 8.0, 1.0])

    ws2 = wb.create_sheet("Склады")
    ws2.append(["Склад", "Артикул WB", "Остатки", "Продажи"])
    ws2.append(["Коледино", 111, 350, 42])
    ws2.append(["Электросталь", 222, 120, 9])

    ws3 = wb.create_sheet("Остатки по размерам")
    ws3.append(["Артикул WB", "Размер", "Остаток"])
    ws3.append([111, "50 мл", 200])
    ws3.append([222, "100 мл", 80])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_workbook_returns_all_sheets():
    sheets = read_workbook(_make_workbook())
    assert set(sheets) == {"Поисковые запросы", "Склады", "Остатки по размерам"}


def test_parse_search_queries():
    sheets = read_workbook(_make_workbook())
    parsed = parse_search_queries_sheet(sheets["Поисковые запросы"])
    assert [a["keyword"] for a in parsed.aggregates] == ["крем для ног", "воск от трещин"]
    assert parsed.aggregates[0]["frequency"] == 5000
    assert parsed.aggregates[1]["frequency_dynamics"] == -30
    assert len(parsed.per_nm) == 3
    first = parsed.per_nm[0]
    assert first["nm_id"] == 111
    assert first["cart_from_search_raw"] == 40
    assert first["order_from_search_raw"] == 7
    assert first["is_rounded_pct"] == 1  # проценты из Excel — округлены WB


def test_parse_warehouse_wide_format():
    sheets = read_workbook(_make_workbook())
    rows = parse_warehouse_sheet(sheets["Склады"])
    assert {(r["warehouse_name"], r["metric_type"]) for r in rows} == {
        ("Коледино", "stock"),
        ("Коледино", "sales"),
        ("Электросталь", "stock"),
        ("Электросталь", "sales"),
    }
    stock_111 = next(
        r for r in rows if r["nm_id"] == 111 and r["metric_type"] == "stock"
    )
    assert stock_111["metric_value"] == 350


def test_parse_warehouse_long_format():
    rows = [
        ["Метрика", "Склад", "Артикул WB", "Значение"],
        ["Остатки", "Тула", 333, 77],
        ["Доставка", "Тула", 333, 5],
    ]
    parsed = parse_warehouse_sheet(rows)
    assert parsed == [
        {"nm_id": 333, "warehouse_name": "Тула", "metric_type": "stock", "metric_value": 77},
        {"nm_id": 333, "warehouse_name": "Тула", "metric_type": "delivery", "metric_value": 5},
    ]


def test_parse_size_stocks():
    sheets = read_workbook(_make_workbook())
    rows = parse_size_stocks_sheet(sheets["Остатки по размерам"])
    assert rows[0]["nm_id"] == 111
    assert rows[0]["size_name"] == "50 мл"
    assert rows[0]["stock_count"] == 200
