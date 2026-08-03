"""Excel ingest — file-manager download → ZIP → XLSX → 4 таблицы.

Excel бесплатен для уже купленных comparison_id. Flow (docs/06-excel-ingest.md):

1. POST /file-manager/download  {"cmpID": ..., "type": "comparison-full"}
2. Poll GET /file-manager/downloads until status == "ready"
3. GET url → ZIP (сохраняем в data/cabinet/excel/)
4. Parse XLSX (calamine, openpyxl fallback) → UPSERT.

Парсеры листов - tolerant header matching (WB меняет заголовки без
предупреждения). Если после WB update ingest даёт 0 rows — смотри реальные
заголовки в сохранённом ZIP и поправь `_match_column` паттерны.
"""

from __future__ import annotations

import asyncio
import io
import json
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
from sqlalchemy.dialects.sqlite import insert as sqlite_insert
from sqlalchemy.ext.asyncio import AsyncEngine

from wb_pool.clients.cabinet import (
    FILE_MANAGER_DOWNLOAD_PATH,
    FILE_MANAGER_DOWNLOADS_PATH,
    WBCabinetClient,
)
from wb_pool.db import UnitOfWork
from wb_pool.logging_setup import get_logger
from wb_pool.models import (
    CmpSearchQuery,
    CmpSearchQueryPerNm,
    CmpSizeStock,
    CmpWarehouseMetric,
)

_log = get_logger(__name__)

_EXCEL_DIR = Path("data/cabinet/excel")
_POLL_ATTEMPTS = 60
_POLL_DELAY_SEC = 2.0


# ---------------------------------------------------------------------------
# Workbook reading (calamine → openpyxl fallback)
# ---------------------------------------------------------------------------


def read_workbook(data: bytes) -> dict[str, list[list[Any]]]:
    """XLSX bytes → {sheet_name: rows}. Calamine first, openpyxl fallback."""
    try:
        from python_calamine import CalamineWorkbook

        wb = CalamineWorkbook.from_filelike(io.BytesIO(data))
        return {
            name: wb.get_sheet_by_name(name).to_python() for name in wb.sheet_names
        }
    except Exception as exc:
        _log.warning("calamine_failed_fallback_openpyxl", error=repr(exc))
        from openpyxl import load_workbook

        wb2 = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out: dict[str, list[list[Any]]] = {}
        for ws in wb2.worksheets:
            out[ws.title] = [[c.value for c in row] for row in ws.iter_rows()]
        return out


# ---------------------------------------------------------------------------
# Pure parsers (tolerant header matching)
# ---------------------------------------------------------------------------


def _norm(value: Any) -> str:
    return str(value or "").strip().lower()


def _match_column(header: list[Any], *needles: str, exclude: tuple[str, ...] = ()) -> int | None:
    """Index первой колонки чей заголовок содержит все needles и ни один exclude."""
    for i, cell in enumerate(header):
        h = _norm(cell)
        if not h:
            continue
        if all(n in h for n in needles) and not any(e in h for e in exclude):
            return i
    return None


def _to_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return round(float(str(value).replace(" ", "").replace(",", ".")))
    except (ValueError, TypeError):
        return None


def _to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(" ", "").replace(",", ".").replace("%", ""))
    except (ValueError, TypeError):
        return None


@dataclass
class SearchQueriesParse:
    """Результат разбора листа «Поисковые запросы»."""

    aggregates: list[dict[str, Any]] = field(default_factory=list)
    per_nm: list[dict[str, Any]] = field(default_factory=list)


def parse_search_queries_sheet(rows: list[list[Any]]) -> SearchQueriesParse:
    """Лист «Поисковые запросы» → aggregate rows + per-nm rows.

    Ожидаемые колонки (tolerant, по подстрокам):
    запрос / артикул / частота / динамика / корзин (переходы) / заказ /
    конверсия в корзину % / конверсия в заказ %.
    """
    out = SearchQueriesParse()
    if not rows:
        return out
    header = rows[0]
    i_kw = _match_column(header, "запрос")
    i_nm = _match_column(header, "артикул")
    i_freq = _match_column(header, "частота", exclude=("динамика",))
    i_freq_dyn = _match_column(header, "динамика")
    i_cart = _match_column(header, "корзин", exclude=("%", "конверси"))
    i_order = _match_column(header, "заказ", exclude=("%", "конверси"))
    i_cart_pct = _match_column(header, "конверси", "корзин")
    i_order_pct = _match_column(header, "конверси", "заказ")
    if i_kw is None:
        return out

    seen_agg: set[str] = set()
    for row in rows[1:]:
        if i_kw >= len(row):
            continue
        kw = str(row[i_kw] or "").strip()
        if not kw:
            continue
        freq = _to_int(row[i_freq]) if i_freq is not None and i_freq < len(row) else None
        freq_dyn = (
            _to_int(row[i_freq_dyn])
            if i_freq_dyn is not None and i_freq_dyn < len(row)
            else None
        )
        if kw not in seen_agg:
            seen_agg.add(kw)
            out.aggregates.append(
                {"keyword": kw, "frequency": freq or 0, "frequency_dynamics": freq_dyn}
            )
        nm = _to_int(row[i_nm]) if i_nm is not None and i_nm < len(row) else None
        if nm is None:
            continue
        cart_pct = (
            _to_float(row[i_cart_pct])
            if i_cart_pct is not None and i_cart_pct < len(row)
            else None
        )
        order_pct = (
            _to_float(row[i_order_pct])
            if i_order_pct is not None and i_order_pct < len(row)
            else None
        )
        # WB округляет проценты до 1% в Excel — точные значения только из
        # raw counts (cart_from_search_raw / open_card_count).
        is_rounded = int(cart_pct is not None or order_pct is not None)
        out.per_nm.append(
            {
                "keyword": kw,
                "nm_id": nm,
                "cart_from_search_raw": (
                    _to_int(row[i_cart]) if i_cart is not None and i_cart < len(row) else None
                ),
                "order_from_search_raw": (
                    _to_int(row[i_order]) if i_order is not None and i_order < len(row) else None
                ),
                "cart_conv_pct_raw": cart_pct,
                "order_conv_pct_raw": order_pct,
                "is_rounded_pct": is_rounded,
            }
        )
    return out


_WAREHOUSE_METRIC_COLUMNS = {
    "остат": "stock",
    "продаж": "sales",
    "достав": "delivery",
}


def parse_warehouse_sheet(rows: list[list[Any]]) -> list[dict[str, Any]]:
    """Лист «Склады» → [{nm_id, warehouse_name, metric_type, metric_value}].

    Поддерживает два формата:
    - long: колонки «Склад / Артикул / Метрика (или Тип) / Значение»
    - wide: колонки «Склад / Артикул / Остатки / Продажи / Доставка»
    """
    if not rows:
        return []
    header = rows[0]
    i_wh = _match_column(header, "склад")
    i_nm = _match_column(header, "артикул")
    if i_wh is None or i_nm is None:
        return []
    i_metric = _match_column(header, "метрика")
    if i_metric is None:
        i_metric = _match_column(header, "тип")
    i_value = _match_column(header, "значен")

    out: list[dict[str, Any]] = []
    if i_metric is not None and i_value is not None:
        for row in rows[1:]:
            nm = _to_int(row[i_nm]) if i_nm < len(row) else None
            value = _to_int(row[i_value]) if i_value < len(row) else None
            if nm is None or value is None:
                continue
            metric_raw = _norm(row[i_metric]) if i_metric < len(row) else ""
            metric = next(
                (v for k, v in _WAREHOUSE_METRIC_COLUMNS.items() if k in metric_raw),
                metric_raw or "stock",
            )
            out.append(
                {
                    "nm_id": nm,
                    "warehouse_name": str(row[i_wh] or "").strip(),
                    "metric_type": metric,
                    "metric_value": value,
                }
            )
        return out

    # wide format
    metric_cols: list[tuple[int, str]] = []
    for i, cell in enumerate(header):
        h = _norm(cell)
        for needle, metric in _WAREHOUSE_METRIC_COLUMNS.items():
            if needle in h:
                metric_cols.append((i, metric))
                break
    for row in rows[1:]:
        nm = _to_int(row[i_nm]) if i_nm < len(row) else None
        if nm is None:
            continue
        wh = str(row[i_wh] or "").strip()
        for i, metric in metric_cols:
            value = _to_int(row[i]) if i < len(row) else None
            if value is None:
                continue
            out.append(
                {
                    "nm_id": nm,
                    "warehouse_name": wh,
                    "metric_type": metric,
                    "metric_value": value,
                }
            )
    return out


def parse_size_stocks_sheet(rows: list[list[Any]]) -> list[dict[str, Any]]:
    """Лист «Остатки по размерам» → [{nm_id, size_name, stock_count, date?}]."""
    if not rows:
        return []
    header = rows[0]
    i_nm = _match_column(header, "артикул")
    i_size = _match_column(header, "размер")
    i_stock = _match_column(header, "остат")
    i_date = _match_column(header, "дата")
    if i_nm is None or i_size is None or i_stock is None:
        return []
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        nm = _to_int(row[i_nm]) if i_nm < len(row) else None
        stock = _to_int(row[i_stock]) if i_stock < len(row) else None
        if nm is None or stock is None:
            continue
        raw_date = row[i_date] if i_date is not None and i_date < len(row) else None
        out.append(
            {
                "nm_id": nm,
                "size_name": str(row[i_size] or "").strip(),
                "stock_count": stock,
                "date_raw": raw_date,
            }
        )
    return out


# ---------------------------------------------------------------------------
# Download + orchestration
# ---------------------------------------------------------------------------


async def _request_and_download_zip(
    client: WBCabinetClient, *, comparison_id: str
) -> tuple[str, bytes]:
    """file-manager flow: request export → poll → download. Returns (id, bytes)."""
    response = await client.post(
        FILE_MANAGER_DOWNLOAD_PATH,
        json={"cmpID": comparison_id, "type": "comparison-full"},
    )
    download_id = str(response.get("downloadId") or response.get("data", {}).get("downloadId"))
    url: str | None = None
    for _ in range(_POLL_ATTEMPTS):
        await asyncio.sleep(_POLL_DELAY_SEC)
        r = await client.get(FILE_MANAGER_DOWNLOADS_PATH)
        downloads = r.get("downloads") or r.get("data", {}).get("downloads") or []
        for d in downloads:
            if str(d.get("downloadId")) == download_id and d.get("status") == "ready":
                url = str(d["url"])
                break
        if url:
            break
    if url is None:
        raise TimeoutError(
            f"Excel export not ready after {int(_POLL_ATTEMPTS * _POLL_DELAY_SEC)} sec"
        )
    async with httpx.AsyncClient(timeout=120.0) as fetcher:
        resp = await fetcher.get(url)
        resp.raise_for_status()
        return download_id, resp.content


def _epoch_midnight_today_utc() -> int:
    now = datetime.now(UTC)
    return int(now.replace(hour=0, minute=0, second=0, microsecond=0).timestamp())


async def import_comparison_to_db(
    client: WBCabinetClient,
    *,
    nm_ids: list[int],
    period_start: Any,
    period_end: Any,
    engine: AsyncEngine,
    ingest_run_id: int,
    comparison_id: str,
) -> tuple[Path, None, dict[str, int]]:
    """Download Excel ZIP для cmp_id, parse, upsert. Idempotent."""
    _ = nm_ids
    download_id, zip_data = await _request_and_download_zip(
        client, comparison_id=comparison_id
    )
    _EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    zip_path = _EXCEL_DIR / f"{download_id}.zip"
    zip_path.write_bytes(zip_data)

    ps_epoch = int(
        datetime.combine(period_start, datetime.min.time(), tzinfo=UTC).timestamp()
    )
    pe_epoch = int(
        datetime.combine(period_end, datetime.min.time(), tzinfo=UTC).timestamp()
    )

    stats = {
        "cmp_search_queries": 0,
        "cmp_search_query_per_nm": 0,
        "cmp_warehouse_metrics": 0,
        "cmp_size_stocks": 0,
    }
    with zipfile.ZipFile(io.BytesIO(zip_data)) as zf:
        for name in zf.namelist():
            if not name.endswith(".xlsx"):
                continue
            sheets = read_workbook(zf.read(name))
            for sheet_name, rows in sheets.items():
                sn = _norm(sheet_name)
                if "поиск" in sn or "запрос" in sn:
                    parsed = parse_search_queries_sheet(rows)
                    written = await _upsert_search_queries(
                        engine,
                        parsed,
                        comparison_id=comparison_id,
                        period_start=ps_epoch,
                        period_end=pe_epoch,
                        ingest_run_id=ingest_run_id,
                    )
                    stats["cmp_search_queries"] += written[0]
                    stats["cmp_search_query_per_nm"] += written[1]
                elif "склад" in sn:
                    wh_rows = parse_warehouse_sheet(rows)
                    stats["cmp_warehouse_metrics"] += await _upsert_warehouse(
                        engine,
                        wh_rows,
                        comparison_id=comparison_id,
                        period_start=ps_epoch,
                        period_end=pe_epoch,
                        ingest_run_id=ingest_run_id,
                    )
                elif "размер" in sn:
                    sz_rows = parse_size_stocks_sheet(rows)
                    stats["cmp_size_stocks"] += await _upsert_size_stocks(
                        engine,
                        sz_rows,
                        comparison_id=comparison_id,
                        ingest_run_id=ingest_run_id,
                    )
    _log.info(
        "excel_ingest_done",
        comparison_id=comparison_id,
        zip=str(zip_path),
        stats=json.dumps(stats),
    )
    return zip_path, None, stats


async def _upsert_search_queries(
    engine: AsyncEngine,
    parsed: SearchQueriesParse,
    *,
    comparison_id: str,
    period_start: int,
    period_end: int,
    ingest_run_id: int,
) -> tuple[int, int]:
    now = int(datetime.now(UTC).timestamp())
    snapshot = _epoch_midnight_today_utc()
    async with UnitOfWork(engine) as uow:
        if parsed.aggregates:
            agg_rows = [
                {
                    "comparison_id": comparison_id,
                    "keyword": a["keyword"],
                    "period_start": period_start,
                    "period_end": period_end,
                    "frequency": a["frequency"],
                    "frequency_dynamics": a["frequency_dynamics"],
                    "snapshot_date": snapshot,
                    "ingested_at": now,
                    "ingest_run_id": ingest_run_id,
                }
                for a in parsed.aggregates
            ]
            stmt = sqlite_insert(CmpSearchQuery).values(agg_rows)
            stmt = stmt.on_conflict_do_update(
                index_elements=[
                    "comparison_id", "keyword", "period_start", "period_end", "snapshot_date"
                ],
                set_={
                    "frequency": stmt.excluded.frequency,
                    "frequency_dynamics": stmt.excluded.frequency_dynamics,
                    "ingested_at": stmt.excluded.ingested_at,
                    "ingest_run_id": stmt.excluded.ingest_run_id,
                },
            )
            await uow.session.execute(stmt)
        if parsed.per_nm:
            nm_rows = [
                {
                    "comparison_id": comparison_id,
                    "keyword": p["keyword"],
                    "nm_id": p["nm_id"],
                    "period_start": period_start,
                    "period_end": period_end,
                    "cart_from_search_raw": p["cart_from_search_raw"],
                    "order_from_search_raw": p["order_from_search_raw"],
                    "cart_conv_pct_raw": p["cart_conv_pct_raw"],
                    "order_conv_pct_raw": p["order_conv_pct_raw"],
                    "is_rounded_pct": p["is_rounded_pct"],
                    "snapshot_date": snapshot,
                    "ingested_at": now,
                    "ingest_run_id": ingest_run_id,
                }
                for p in parsed.per_nm
            ]
            stmt2 = sqlite_insert(CmpSearchQueryPerNm).values(nm_rows)
            stmt2 = stmt2.on_conflict_do_update(
                index_elements=[
                    "comparison_id", "keyword", "nm_id",
                    "period_start", "period_end", "snapshot_date",
                ],
                set_={
                    "cart_from_search_raw": stmt2.excluded.cart_from_search_raw,
                    "order_from_search_raw": stmt2.excluded.order_from_search_raw,
                    "cart_conv_pct_raw": stmt2.excluded.cart_conv_pct_raw,
                    "order_conv_pct_raw": stmt2.excluded.order_conv_pct_raw,
                    "is_rounded_pct": stmt2.excluded.is_rounded_pct,
                    "ingested_at": stmt2.excluded.ingested_at,
                    "ingest_run_id": stmt2.excluded.ingest_run_id,
                },
            )
            await uow.session.execute(stmt2)
        await uow.commit()
    return len(parsed.aggregates), len(parsed.per_nm)


async def _upsert_warehouse(
    engine: AsyncEngine,
    rows: list[dict[str, Any]],
    *,
    comparison_id: str,
    period_start: int,
    period_end: int,
    ingest_run_id: int,
) -> int:
    if not rows:
        return 0
    snapshot = _epoch_midnight_today_utc()
    values = [
        {
            "comparison_id": comparison_id,
            "nm_id": r["nm_id"],
            "warehouse_name": r["warehouse_name"],
            "metric_type": r["metric_type"],
            "metric_value": r["metric_value"],
            "period_start": period_start,
            "period_end": period_end,
            "snapshot_date": snapshot,
            "ingest_run_id": ingest_run_id,
        }
        for r in rows
    ]
    async with UnitOfWork(engine) as uow:
        stmt = sqlite_insert(CmpWarehouseMetric).values(values)
        stmt = stmt.on_conflict_do_update(
            index_elements=[
                "comparison_id", "nm_id", "warehouse_name", "metric_type",
                "period_start", "period_end", "snapshot_date",
            ],
            set_={
                "metric_value": stmt.excluded.metric_value,
                "ingest_run_id": stmt.excluded.ingest_run_id,
            },
        )
        await uow.session.execute(stmt)
        await uow.commit()
    return len(values)


async def _upsert_size_stocks(
    engine: AsyncEngine,
    rows: list[dict[str, Any]],
    *,
    comparison_id: str,
    ingest_run_id: int,
) -> int:
    if not rows:
        return 0
    snapshot = _epoch_midnight_today_utc()
    values = []
    for r in rows:
        raw_date = r.get("date_raw")
        if isinstance(raw_date, datetime):
            date_epoch = int(raw_date.replace(tzinfo=UTC).timestamp())
        elif isinstance(raw_date, str) and len(raw_date) >= 10:
            try:
                date_epoch = int(
                    datetime.strptime(raw_date[:10], "%Y-%m-%d")
                    .replace(tzinfo=UTC)
                    .timestamp()
                )
            except ValueError:
                date_epoch = snapshot
        else:
            date_epoch = snapshot
        values.append(
            {
                "comparison_id": comparison_id,
                "nm_id": r["nm_id"],
                "size_name": r["size_name"],
                "stock_count": r["stock_count"],
                "date": date_epoch,
                "ingest_run_id": ingest_run_id,
            }
        )
    async with UnitOfWork(engine) as uow:
        stmt = sqlite_insert(CmpSizeStock).values(values)
        stmt = stmt.on_conflict_do_update(
            index_elements=["comparison_id", "nm_id", "size_name", "date"],
            set_={
                "stock_count": stmt.excluded.stock_count,
                "ingest_run_id": stmt.excluded.ingest_run_id,
            },
        )
        await uow.session.execute(stmt)
        await uow.commit()
    return len(values)


__all__ = [
    "SearchQueriesParse",
    "import_comparison_to_db",
    "parse_search_queries_sheet",
    "parse_size_stocks_sheet",
    "parse_warehouse_sheet",
    "read_workbook",
]
